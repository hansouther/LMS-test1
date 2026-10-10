from fastapi import APIRouter, HTTPException, Depends, BackgroundTasks, UploadFile, File
from fastapi.responses import Response
from pydantic import BaseModel, EmailStr, Field
from typing import List, Optional
import io
import csv
import re
import zipfile
import asyncio
import openpyxl

from database import db
from utils import new_id, now_iso
from security import require_roles, hash_password
from emailer import notify_new_tryout, notify_bid_accepted, notify_new_material
from storage import put_object, MIME_TYPES, APP_NAME
from routes_files import save_question_image, IMAGE_EXT
import notifications
import analysis
import irt
from grading import grade_attempt
from tryout_session import schedule_status, remaining_seconds, answered_count

router = APIRouter(prefix="/api/admin", tags=["admin"])
admin_only = require_roles("admin")


async def _active_student_recipients():
    enrolls = await db.enrollments.find({}, {"_id": 0, "student_id": 1}).to_list(5000)
    ids = list({e["student_id"] for e in enrolls})
    students = await db.users.find({"id": {"$in": ids}, "role": "student"}, {"_id": 0, "email": 1, "name": 1}).to_list(5000)
    return [{"email": s["email"], "name": s.get("name")} for s in students if s.get("email")]


# ---------- Models ----------
class NewsBody(BaseModel):
    title: str
    content: str
    category: str = "Pengumuman"
    published: bool = True


class CalendarBody(BaseModel):
    title: str
    description: Optional[str] = None
    date: str
    type: str = "academic"


class CourseBody(BaseModel):
    title: str
    description: str
    subject: str
    level: str = "Umum"
    price: int = 0
    thumbnail: Optional[str] = None
    active: bool = True


class SessionBody(BaseModel):
    id: Optional[str] = None
    date: str
    start_time: str
    end_time: str
    topic: Optional[str] = None


class SlotBody(BaseModel):
    title: str
    subject: str
    required_qualifications: List[str] = []
    course_id: Optional[str] = None
    open_to_all: bool = False
    notes: Optional[str] = None
    sessions: List[SessionBody] = []
    # legacy single-session support
    date: Optional[str] = None
    start_time: Optional[str] = None
    end_time: Optional[str] = None


class TryoutBody(BaseModel):
    title: str
    description: Optional[str] = None
    subject: str
    duration_minutes: int = 60
    start_at: Optional[str] = None
    end_at: Optional[str] = None
    published: bool = False
    course_id: Optional[str] = None
    kind: str = "standalone"
    irt_scale: str = irt.DEFAULT_SCALE  # snbt | tka | raw | custom
    irt_scale_custom: Optional[dict] = None  # {min, max, mean, sd} bila custom
    price: int = 0  # harga Try Out (Rupiah); 0 = gratis
    shuffle_questions: bool = False
    shuffle_options: bool = False
    tab_warning_limit: int = Field(5, ge=0, le=20)  # 0 = deteksi pindah tab nonaktif
    show_review: bool = False  # pembahasan & kunci dibuka untuk siswa


class IrtScaleBody(BaseModel):
    irt_scale: str
    irt_scale_custom: Optional[dict] = None


def _clean_scale(key, custom):
    key = (key or irt.DEFAULT_SCALE).lower()
    if key not in irt.SCALE_PRESETS:
        raise HTTPException(status_code=400, detail="Skala skor IRT tidak valid")
    if key == "custom":
        custom, err = irt.validate_custom_scale(custom or {})
        if err:
            raise HTTPException(status_code=400, detail=err)
        return key, custom
    return key, None


async def _rescore_attempts(tryout_id: str, tryout: dict):
    """Hitung ulang theta/skor IRT semua attempt tersubmit memakai skala Try Out."""
    scale = irt.resolve_scale(tryout)
    questions = await db.questions.find({"tryout_id": tryout_id}, {"_id": 0}).to_list(1000)
    attempts = await db.attempts.find({"tryout_id": tryout_id, "status": "submitted"}, {"_id": 0}).to_list(10000)
    for a in attempts:
        res = irt.score_attempt(questions, a.get("per_question"), scale)
        await db.attempts.update_one({"id": a["id"]}, {"$set": res})
    from routes_student import align_tryout_irt
    await align_tryout_irt(tryout_id)
    return len(attempts)


class QuestionBody(BaseModel):
    type: str  # single | multiple | truefalse | truefalse_multi | essay
    text: str
    image_url: Optional[str] = None
    tf_labels: Optional[List[str]] = None
    options: List[dict] = []
    correct_answers: List[str] = []
    points: int = 1
    order: int = 0
    competency: str = "umum"  # numerasi | literasi | umum (AKM)
    difficulty_label: Optional[str] = None  # mudah | sedang | sulit (label manual untuk IRT)
    bidang: Optional[str] = None  # mata pelajaran
    materi: Optional[str] = None  # topik
    explanation: Optional[str] = None  # pembahasan (mendukung rumus)
    explanation_image_url: Optional[str] = None


class CopyQuestionsBody(BaseModel):
    question_ids: List[str]


class BulkMateriBody(BaseModel):
    question_ids: List[str]
    bidang: Optional[str] = None
    materi: Optional[str] = None


def _clean_text(v):
    return (v or "").strip() or None


async def backfill_question_bidang():
    async for t in db.tryouts.find({}, {"_id": 0, "id": 1, "subject": 1}):
        if t.get("subject"):
            await db.questions.update_many({"tryout_id": t["id"], "$or": [{"bidang": None}, {"bidang": {"$exists": False}}, {"bidang": ""}]},
                                           {"$set": {"bidang": t["subject"].strip()}})


class BulkCompetencyBody(BaseModel):
    question_ids: List[str]
    competency: str  # numerasi | literasi | umum


class BulkDifficultyBody(BaseModel):
    question_ids: List[str]
    difficulty_label: str  # mudah | sedang | sulit


class BroadcastBody(BaseModel):
    title: str
    message: str
    priority: str = "normal"


class UserBody(BaseModel):
    name: str
    email: EmailStr
    password: str
    role: str  # tutor | proctor | student
    phone: Optional[str] = None
    school_id: Optional[str] = None
    qualifications: List[str] = []


class UserUpdateBody(BaseModel):
    role: Optional[str] = None
    school_id: Optional[str] = None
    status: Optional[str] = None  # pending | approved | rejected


class SchoolBody(BaseModel):
    name: str
    city: Optional[str] = None


# ---------- Dashboard ----------
@router.get("/stats")
async def stats(user: dict = Depends(admin_only)):
    return {
        "students": await db.users.count_documents({"role": "student"}),
        "tutors": await db.users.count_documents({"role": "tutor"}),
        "proctors": await db.users.count_documents({"role": "proctor"}),
        "courses": await db.courses.count_documents({}),
        "tryouts": await db.tryouts.count_documents({}),
        "open_slots": await db.teaching_slots.count_documents({"status": "open"}),
        "pending_partnerships": await db.partnerships.count_documents({"status": "new"}),
        "pending_verifications": await db.users.count_documents({"status": "pending"}),
        "total_attempts": await db.attempts.count_documents({"status": "submitted"}),
    }


@router.get("/pending-count")
async def pending_count(user: dict = Depends(admin_only)):
    return {"count": await db.users.count_documents({"status": "pending"})}


# ---------- News ----------
@router.get("/news")
async def list_news(user: dict = Depends(admin_only)):
    return await db.news.find({}, {"_id": 0}).sort("created_at", -1).to_list(200)


@router.post("/news")
async def create_news(body: NewsBody, user: dict = Depends(admin_only)):
    doc = {"id": new_id(), **body.model_dump(), "author": user["name"], "created_at": now_iso()}
    await db.news.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


@router.put("/news/{item_id}")
async def update_news(item_id: str, body: NewsBody, user: dict = Depends(admin_only)):
    res = await db.news.update_one({"id": item_id}, {"$set": body.model_dump()})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Berita tidak ditemukan")
    return {"ok": True}


@router.delete("/news/{item_id}")
async def delete_news(item_id: str, user: dict = Depends(admin_only)):
    await db.news.delete_one({"id": item_id})
    return {"ok": True}


# ---------- Calendar ----------
@router.get("/calendar")
async def list_calendar(user: dict = Depends(admin_only)):
    return await db.calendar_events.find({}, {"_id": 0}).sort("date", 1).to_list(500)


@router.post("/calendar")
async def create_calendar(body: CalendarBody, user: dict = Depends(admin_only)):
    doc = {"id": new_id(), **body.model_dump(), "created_at": now_iso()}
    await db.calendar_events.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


@router.delete("/calendar/{item_id}")
async def delete_calendar(item_id: str, user: dict = Depends(admin_only)):
    await db.calendar_events.delete_one({"id": item_id})
    return {"ok": True}


# ---------- Courses ----------
@router.get("/courses")
async def list_courses(user: dict = Depends(admin_only)):
    items = await db.courses.find({}, {"_id": 0}).sort("created_at", -1).to_list(200)
    for c in items:
        c["enrolled_count"] = await db.enrollments.count_documents({"course_id": c["id"]})
    return items


@router.post("/courses")
async def create_course(body: CourseBody, user: dict = Depends(admin_only)):
    doc = {"id": new_id(), **body.model_dump(), "created_at": now_iso()}
    await db.courses.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


@router.put("/courses/{item_id}")
async def update_course(item_id: str, body: CourseBody, user: dict = Depends(admin_only)):
    res = await db.courses.update_one({"id": item_id}, {"$set": body.model_dump()})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Kursus tidak ditemukan")
    return {"ok": True}


@router.delete("/courses/{item_id}")
async def delete_course(item_id: str, user: dict = Depends(admin_only)):
    await db.courses.delete_one({"id": item_id})
    return {"ok": True}


# ---------- Partnerships ----------
@router.get("/partnerships")
async def list_partnerships(user: dict = Depends(admin_only)):
    return await db.partnerships.find({}, {"_id": 0}).sort("created_at", -1).to_list(200)


@router.put("/partnerships/{item_id}")
async def update_partnership(item_id: str, status: str, user: dict = Depends(admin_only)):
    await db.partnerships.update_one({"id": item_id}, {"$set": {"status": status}})
    return {"ok": True}


# ---------- Teaching Slots & Bidding ----------
@router.get("/slots")
async def list_slots(user: dict = Depends(admin_only)):
    slots = await db.teaching_slots.find({}, {"_id": 0}).sort("date", 1).to_list(300)
    for s in slots:
        s["bid_count"] = await db.bids.count_documents({"slot_id": s["id"]})
    return slots


def _build_sessions(body):
    sessions = []
    for i, s in enumerate(body.sessions):
        sessions.append({
            "id": s.id or new_id(),
            "no": i + 1,
            "date": s.date,
            "start_time": s.start_time,
            "end_time": s.end_time,
            "topic": s.topic,
        })
    if not sessions and body.date:
        sessions = [{
            "id": new_id(), "no": 1, "date": body.date,
            "start_time": body.start_time, "end_time": body.end_time, "topic": body.notes,
        }]
    return sessions


@router.post("/slots")
async def create_slot(body: SlotBody, user: dict = Depends(admin_only)):
    sessions = _build_sessions(body)
    if not sessions:
        raise HTTPException(status_code=400, detail="Minimal satu pertemuan diperlukan")
    first = sessions[0]
    doc = {
        "id": new_id(),
        "title": body.title,
        "subject": body.subject,
        "required_qualifications": body.required_qualifications,
        "course_id": body.course_id,
        "open_to_all": body.open_to_all,
        "notes": body.notes,
        "sessions": sessions,
        "date": first["date"],
        "start_time": first["start_time"],
        "end_time": first["end_time"],
        "status": "open",
        "tutor_id": None,
        "created_by": user["id"],
        "created_at": now_iso(),
    }
    await db.teaching_slots.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


@router.put("/slots/{slot_id}")
async def update_slot(slot_id: str, body: SlotBody, user: dict = Depends(admin_only)):
    slot = await db.teaching_slots.find_one({"id": slot_id})
    if not slot:
        raise HTTPException(status_code=404, detail="Slot tidak ditemukan")
    sessions = _build_sessions(body)
    if not sessions:
        raise HTTPException(status_code=400, detail="Minimal satu pertemuan diperlukan")
    first = sessions[0]
    old_count = len(slot.get("sessions") or [])
    updates = {
        "title": body.title, "subject": body.subject,
        "required_qualifications": body.required_qualifications,
        "course_id": body.course_id, "open_to_all": body.open_to_all, "notes": body.notes,
        "sessions": sessions, "date": first["date"],
        "start_time": first["start_time"], "end_time": first["end_time"],
    }
    await db.teaching_slots.update_one({"id": slot_id}, {"$set": updates})
    # Notify students if new sessions were added to a confirmed class
    if slot.get("status") == "confirmed" and len(sessions) > old_count:
        added = len(sessions) - old_count
        sids = await notifications.course_student_ids(body.course_id)
        if sids:
            await notifications.push(
                sids, "session",
                f"{added} pertemuan baru ditambahkan",
                f'Kelas "{body.title}" kini memiliki {len(sessions)} pertemuan',
                "/student/schedule",
            )
            students = await db.users.find({"id": {"$in": sids}}, {"_id": 0, "email": 1, "name": 1}).to_list(2000)
            await notify_new_material(students, body.title, f"{added} pertemuan baru ditambahkan ke jadwal kelas")
    return {"ok": True}


@router.delete("/slots/{slot_id}")
async def delete_slot(slot_id: str, user: dict = Depends(admin_only)):
    await db.teaching_slots.delete_one({"id": slot_id})
    await db.bids.delete_many({"slot_id": slot_id})
    await db.class_materials.delete_many({"slot_id": slot_id})
    await db.attendance.delete_many({"slot_id": slot_id})
    return {"ok": True}


@router.get("/slots/{slot_id}/bids")
async def slot_bids(slot_id: str, user: dict = Depends(admin_only)):
    return await db.bids.find({"slot_id": slot_id}, {"_id": 0}).sort("created_at", 1).to_list(200)


@router.post("/slots/{slot_id}/assign/{bid_id}")
async def assign_slot(slot_id: str, bid_id: str, background: BackgroundTasks, user: dict = Depends(admin_only)):
    bid = await db.bids.find_one({"id": bid_id, "slot_id": slot_id}, {"_id": 0})
    if not bid:
        raise HTTPException(status_code=404, detail="Bid tidak ditemukan")
    await db.teaching_slots.update_one(
        {"id": slot_id}, {"$set": {"status": "confirmed", "tutor_id": bid["tutor_id"]}}
    )
    await db.bids.update_one({"id": bid_id}, {"$set": {"status": "accepted"}})
    await db.bids.update_many(
        {"slot_id": slot_id, "id": {"$ne": bid_id}}, {"$set": {"status": "rejected"}}
    )
    slot = await db.teaching_slots.find_one({"id": slot_id}, {"_id": 0})
    tutor = await db.users.find_one({"id": bid["tutor_id"]}, {"_id": 0, "email": 1, "name": 1})
    if slot and tutor and tutor.get("email"):
        background.add_task(
            notify_bid_accepted, tutor["email"], tutor.get("name", "Tentor"),
            slot.get("title", ""), slot.get("date", ""),
            f'{slot.get("start_time", "")}-{slot.get("end_time", "")}',
        )
    return {"ok": True}


# ---------- Try Out & Questions ----------
@router.get("/tryouts")
async def list_tryouts(user: dict = Depends(admin_only)):
    items = await db.tryouts.find({"kind": {"$ne": "exercise"}}, {"_id": 0}).sort("created_at", -1).to_list(200)
    for t in items:
        t["question_count"] = await db.questions.count_documents({"tryout_id": t["id"]})
        t["attempt_count"] = await db.attempts.count_documents({"tryout_id": t["id"], "status": "submitted"})
        t["in_progress_count"] = await db.attempts.count_documents({"tryout_id": t["id"], "status": "in_progress"})
        t["schedule_status"] = schedule_status(t)
    return items


@router.post("/tryouts")
async def create_tryout(body: TryoutBody, background: BackgroundTasks, user: dict = Depends(admin_only)):
    key, custom = _clean_scale(body.irt_scale, body.irt_scale_custom)
    doc = {"id": new_id(), **body.model_dump(), "irt_scale": key, "irt_scale_custom": custom,
           "created_by": user["id"], "created_at": now_iso()}
    await db.tryouts.insert_one(doc)
    if doc.get("published") and doc.get("kind", "standalone") == "standalone":
        recipients = await _active_student_recipients()
        if recipients:
            background.add_task(notify_new_tryout, recipients, doc["title"], doc["subject"])
    return {k: v for k, v in doc.items() if k != "_id"}


@router.put("/tryouts/{tryout_id}")
async def update_tryout(tryout_id: str, body: TryoutBody, background: BackgroundTasks, user: dict = Depends(admin_only)):
    existing = await db.tryouts.find_one({"id": tryout_id}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    key, custom = _clean_scale(body.irt_scale, body.irt_scale_custom)
    updates = {**body.model_dump(), "irt_scale": key, "irt_scale_custom": custom}
    await db.tryouts.update_one({"id": tryout_id}, {"$set": updates})
    scale_changed = (existing.get("irt_scale") or irt.DEFAULT_SCALE) != key or (existing.get("irt_scale_custom") or None) != custom
    rescored = await _rescore_attempts(tryout_id, {**existing, **updates}) if scale_changed else 0
    if body.published and not existing.get("published") and body.kind == "standalone":
        recipients = await _active_student_recipients()
        if recipients:
            background.add_task(notify_new_tryout, recipients, body.title, body.subject)
    return {"ok": True, "rescored": rescored}


class ReviewBody(BaseModel):
    show_review: bool


@router.put("/tryouts/{tryout_id}/review")
async def set_review(tryout_id: str, body: ReviewBody, user: dict = Depends(admin_only)):
    res = await db.tryouts.update_one({"id": tryout_id}, {"$set": {"show_review": body.show_review}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    return {"ok": True, "show_review": body.show_review}


@router.get("/irt-scales")
async def irt_scales(user: dict = Depends(admin_only)):
    """Daftar preset skala skor IRT (SNBT, TKA, SD/SMP, Kustom)."""
    return list(irt.SCALE_PRESETS.values())


@router.get("/tryouts/{tryout_id}/irt-scale")
async def get_irt_scale(tryout_id: str, user: dict = Depends(admin_only)):
    t = await db.tryouts.find_one({"id": tryout_id}, {"_id": 0})
    if not t:
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    return {"irt_scale": t.get("irt_scale") or irt.DEFAULT_SCALE,
            "irt_scale_custom": t.get("irt_scale_custom"),
            "effective": irt.resolve_scale(t)}


@router.put("/tryouts/{tryout_id}/irt-scale")
async def set_irt_scale(tryout_id: str, body: IrtScaleBody, user: dict = Depends(admin_only)):
    """Ubah skala skor IRT sebuah Try Out lalu hitung ulang skor semua attempt tersubmit."""
    t = await db.tryouts.find_one({"id": tryout_id}, {"_id": 0})
    if not t:
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    key, custom = _clean_scale(body.irt_scale, body.irt_scale_custom)
    await db.tryouts.update_one({"id": tryout_id}, {"$set": {"irt_scale": key, "irt_scale_custom": custom}})
    t.update({"irt_scale": key, "irt_scale_custom": custom})
    rescored = await _rescore_attempts(tryout_id, t)
    return {"ok": True, "irt_scale": key, "irt_scale_custom": custom,
            "effective": irt.resolve_scale(t), "rescored": rescored}


@router.delete("/tryouts/{tryout_id}")
async def delete_tryout(tryout_id: str, user: dict = Depends(admin_only)):
    await db.tryouts.delete_one({"id": tryout_id})
    await db.questions.delete_many({"tryout_id": tryout_id})
    await db.tryout_violations.delete_many({"tryout_id": tryout_id})
    return {"ok": True}


@router.get("/tryouts/{tryout_id}/questions")
async def list_questions(tryout_id: str, user: dict = Depends(admin_only)):
    return await db.questions.find({"tryout_id": tryout_id}, {"_id": 0}).sort("order", 1).to_list(300)


@router.post("/tryouts/{tryout_id}/questions")
async def create_question(tryout_id: str, body: QuestionBody, user: dict = Depends(admin_only)):
    tryout = await db.tryouts.find_one({"id": tryout_id}, {"_id": 0, "subject": 1})
    if not tryout:
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    count = await db.questions.count_documents({"tryout_id": tryout_id})
    payload = body.model_dump()
    payload["bidang"] = _clean_text(payload.get("bidang")) or _clean_text(tryout.get("subject"))
    payload["materi"] = _clean_text(payload.get("materi"))
    dl = (payload.get("difficulty_label") or "").lower()
    payload["difficulty_label"] = dl if dl in irt.DIFFICULTY_LABELS else None
    doc = {"id": new_id(), "tryout_id": tryout_id, **payload,
           "irt_a": None, "irt_b": None, "irt_c": None,
           "irt_calibrated": False, "p_value": None, "point_biserial": None,
           "response_count": 0}
    if not doc.get("order"):
        doc["order"] = count + 1
    await db.questions.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


def _bank_query(exclude, bidang, materi, competency, q):
    query = {"tryout_id": {"$ne": exclude}} if exclude else {}
    if bidang:
        query["bidang"] = None if bidang == "Tanpa bidang" else bidang
    if materi:
        query["materi"] = materi
    if competency:
        query["competency"] = competency
    if q:
        query["text"] = {"$regex": re.escape(q.strip()), "$options": "i"}
    return query


@router.get("/question-bank")
async def question_bank(exclude: Optional[str] = None, bidang: Optional[str] = None, materi: Optional[str] = None,
                        competency: Optional[str] = None, difficulty: Optional[str] = None, q: Optional[str] = None,
                        user: dict = Depends(admin_only)):
    base = {"tryout_id": {"$ne": exclude}} if exclude else {}
    facets_rows = await db.questions.find(base, {"_id": 0, "bidang": 1, "materi": 1}).to_list(50000)
    materi_by_bidang = {}
    for r in facets_rows:
        b = r.get("bidang") or "Tanpa bidang"
        materi_by_bidang.setdefault(b, set())
        if r.get("materi"):
            materi_by_bidang[b].add(r["materi"])
    proj = {"_id": 0, "id": 1, "tryout_id": 1, "type": 1, "text": 1, "points": 1, "competency": 1, "bidang": 1,
            "materi": 1, "difficulty_label": 1, "irt_calibrated": 1, "irt_b": 1, "image_url": 1}
    items = await db.questions.find(_bank_query(exclude, bidang, materi, competency, q), proj).sort([("bidang", 1), ("materi", 1)]).to_list(3000)
    for it in items:
        it["difficulty"] = irt.effective_label(it)
    if difficulty:
        items = [it for it in items if it["difficulty"] == difficulty]
    titles = {t["id"]: t["title"] for t in await db.tryouts.find(
        {"id": {"$in": list({it["tryout_id"] for it in items})}}, {"_id": 0, "id": 1, "title": 1}).to_list(5000)}
    for it in items:
        it["source_title"] = titles.get(it["tryout_id"])
    return {"items": items, "facets": {"bidang": sorted(materi_by_bidang),
                                       "materi": {k: sorted(v) for k, v in materi_by_bidang.items()}}}


def _question_to_row(q: dict) -> dict:
    t = q.get("type")
    opts = q.get("options") or []
    correct = q.get("correct_answers") or []
    row = {"type": t, "text": q.get("text", ""), "points": q.get("points", 1), "competency": q.get("competency") or "umum",
           "tingkat": irt.effective_label(q), "bidang": q.get("bidang") or "", "materi": q.get("materi") or "",
           "pembahasan": q.get("explanation") or "", "gambar": q.get("image_url") or "",
           "gambar_pembahasan": q.get("explanation_image_url") or ""}
    if t in ("single", "multiple", "truefalse_multi"):
        for i, o in enumerate(opts[:len(_LETTERS)]):
            row[f"option_{_LETTERS[i]}"] = o.get("text", "")
    idx = {o.get("id"): i for i, o in enumerate(opts)}
    if t in ("single", "multiple"):
        row["correct"] = ";".join(_LETTERS[idx[c]].upper() for c in correct if c in idx and idx[c] < len(_LETTERS))
    elif t == "truefalse":
        row["correct"] = "benar" if (correct[:1] or [""])[0] == "true" else "salah"
    elif t == "truefalse_multi":
        key = dict(c.split(":", 1) for c in correct if ":" in c)
        row["correct"] = ";".join("B" if key.get(o.get("id")) == "true" else "S" for o in opts[:len(_LETTERS)])
    else:
        row["correct"] = "|".join(correct)
    if t in ("truefalse", "truefalse_multi") and q.get("tf_labels"):
        row["label_benar"], row["label_salah"] = (q["tf_labels"] + ["", ""])[:2]
    return row


@router.get("/question-bank/export")
async def question_bank_export(bidang: Optional[str] = None, materi: Optional[str] = None,
                               competency: Optional[str] = None, difficulty: Optional[str] = None,
                               q: Optional[str] = None, user: dict = Depends(admin_only)):
    items = await db.questions.find(_bank_query(None, bidang, materi, competency, q), {"_id": 0}).sort(
        [("bidang", 1), ("materi", 1), ("tryout_id", 1), ("order", 1)]).to_list(20000)
    if difficulty:
        items = [it for it in items if irt.effective_label(it) == difficulty]
    titles = {t["id"]: t["title"] for t in await db.tryouts.find(
        {"id": {"$in": list({it["tryout_id"] for it in items})}}, {"_id": 0, "id": 1, "title": 1}).to_list(5000)}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Soal"
    headers = _TPL_HEADERS + ["sumber_tryout"]
    ws.append(headers)
    for it in items:
        row = {**_question_to_row(it), "sumber_tryout": titles.get(it["tryout_id"], "")}
        ws.append([row.get(h, "") for h in headers])
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(max(width, 10), 48)
    info = wb.create_sheet("Info")
    filt = {"bidang": bidang, "materi": materi, "kompetensi": competency, "tingkat": difficulty, "cari": q}
    info.append(["Ekspor Bank Soal", now_iso()])
    info.append(["Jumlah soal", len(items)])
    for k, v in filt.items():
        info.append([k, v or "Semua"])
    info.append(["Catatan", "Format sama dengan template impor; berkas ini bisa langsung diimpor ke Try Out lain. Kolom sumber_tryout diabaikan saat impor."])
    buf = io.BytesIO()
    wb.save(buf)
    slug = re.sub(r"[^a-z0-9]+", "_", "_".join(x for x in (bidang, materi) if x).lower()).strip("_") or "semua"
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=bank_soal_{slug}.xlsx"})


@router.post("/tryouts/{tryout_id}/questions/bulk-materi")
async def bulk_materi(tryout_id: str, body: BulkMateriBody, user: dict = Depends(admin_only)):
    upd = {k: _clean_text(getattr(body, k)) for k in ("bidang", "materi") if getattr(body, k) is not None}
    if not upd or not body.question_ids:
        raise HTTPException(status_code=400, detail="Isi bidang atau materi dan pilih soal")
    res = await db.questions.update_many({"tryout_id": tryout_id, "id": {"$in": body.question_ids}}, {"$set": upd})
    return {"updated": res.modified_count}


@router.post("/tryouts/{tryout_id}/questions/copy")
async def copy_questions(tryout_id: str, body: CopyQuestionsBody, user: dict = Depends(admin_only)):
    if not await db.tryouts.find_one({"id": tryout_id}):
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    src = await db.questions.find({"id": {"$in": body.question_ids}, "tryout_id": {"$ne": tryout_id}}, {"_id": 0}).sort("order", 1).to_list(5000)
    count = await db.questions.count_documents({"tryout_id": tryout_id})
    docs = []
    for i, q in enumerate(src):
        docs.append({**q, "id": new_id(), "tryout_id": tryout_id, "order": count + i + 1, "copied_from": q["id"],
                     "irt_a": None, "irt_b": None, "irt_c": None, "irt_calibrated": False,
                     "p_value": None, "point_biserial": None, "response_count": 0})
    if docs:
        await db.questions.insert_many(docs)
    return {"copied": len(docs)}


@router.put("/questions/{question_id}")
async def update_question(question_id: str, body: QuestionBody, user: dict = Depends(admin_only)):
    payload = body.model_dump()
    dl = (payload.get("difficulty_label") or "").lower()
    payload["difficulty_label"] = dl if dl in irt.DIFFICULTY_LABELS else None
    payload["bidang"] = _clean_text(payload.get("bidang"))
    payload["materi"] = _clean_text(payload.get("materi"))
    res = await db.questions.update_one({"id": question_id}, {"$set": payload})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Soal tidak ditemukan")
    return {"ok": True}


@router.delete("/questions/{question_id}")
async def delete_question(question_id: str, user: dict = Depends(admin_only)):
    await db.questions.delete_one({"id": question_id})
    return {"ok": True}


@router.post("/tryouts/{tryout_id}/questions/bulk-competency")
async def bulk_competency(tryout_id: str, body: BulkCompetencyBody, user: dict = Depends(admin_only)):
    comp = body.competency if body.competency in ("numerasi", "literasi", "umum") else "umum"
    if not body.question_ids:
        raise HTTPException(status_code=400, detail="Tidak ada soal dipilih")
    res = await db.questions.update_many(
        {"tryout_id": tryout_id, "id": {"$in": body.question_ids}},
        {"$set": {"competency": comp}},
    )
    return {"updated": res.modified_count}


@router.post("/tryouts/{tryout_id}/questions/bulk-difficulty")
async def bulk_difficulty(tryout_id: str, body: BulkDifficultyBody, user: dict = Depends(admin_only)):
    lvl = (body.difficulty_label or "").lower()
    if lvl not in irt.DIFFICULTY_LABELS:
        raise HTTPException(status_code=400, detail="Tingkat kesulitan tidak valid")
    if not body.question_ids:
        raise HTTPException(status_code=400, detail="Tidak ada soal dipilih")
    res = await db.questions.update_many(
        {"tryout_id": tryout_id, "id": {"$in": body.question_ids}},
        {"$set": {"difficulty_label": lvl}},
    )
    return {"updated": res.modified_count}


@router.post("/tryouts/{tryout_id}/calibrate")
async def calibrate_tryout(tryout_id: str, user: dict = Depends(admin_only)):
    """Kalibrasi parameter IRT (3PL) tiap butir dari data jawaban peserta,
    lalu hitung ulang skor IR/theta untuk semua attempt tersubmit."""
    if not await db.tryouts.find_one({"id": tryout_id}):
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    tryout = await db.tryouts.find_one({"id": tryout_id}, {"_id": 0})
    scale = irt.resolve_scale(tryout)
    questions = await db.questions.find({"tryout_id": tryout_id}, {"_id": 0}).to_list(1000)
    attempts = await db.attempts.find(
        {"tryout_id": tryout_id, "status": "submitted"}, {"_id": 0}
    ).to_list(10000)

    params = irt.calibrate(questions, attempts)
    calibrated = 0
    for qid, p in params.items():
        await db.questions.update_one({"id": qid}, {"$set": p})
        if p.get("irt_calibrated"):
            calibrated += 1

    # Terapkan parameter baru ke objek soal in-memory lalu skor ulang theta tiap attempt
    qmap = {q["id"]: q for q in questions}
    for qid, p in params.items():
        if qid in qmap:
            qmap[qid].update(p)
    rescored = 0
    for a in attempts:
        res = irt.score_attempt(list(qmap.values()), a.get("per_question"), scale)
        await db.attempts.update_one({"id": a["id"]}, {"$set": res})
        rescored += 1
    from routes_student import align_tryout_irt
    await align_tryout_irt(tryout_id)

    return {
        "calibrated": calibrated,
        "items_with_responses": len(params),
        "total_questions": len(questions),
        "attempts_scored": rescored,
        "min_responses_required": 5,
    }


@router.get("/tryouts/{tryout_id}/irt-blueprint")
async def irt_blueprint(tryout_id: str, user: dict = Depends(admin_only)):
    """Rekomendasi komposisi soal (Mudah/Sedang/Sulit) + Test Information Function
    untuk memenuhi kebutuhan penilaian IRT saat membuat Try Out."""
    if not await db.tryouts.find_one({"id": tryout_id}):
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    questions = await db.questions.find({"tryout_id": tryout_id}, {"_id": 0}).to_list(1000)
    return irt.build_blueprint(questions)


@router.get("/tryouts/{tryout_id}/results")
async def tryout_results(tryout_id: str, user: dict = Depends(admin_only)):
    attempts = await db.attempts.find(
        {"tryout_id": tryout_id, "status": "submitted"}, {"_id": 0, "per_question": 0}
    ).to_list(2000)
    sids = list({a["student_id"] for a in attempts})
    students = await db.users.find({"id": {"$in": sids}}, {"_id": 0, "id": 1, "name": 1, "email": 1}).to_list(2000)
    smap = {s["id"]: s for s in students}
    cheat = await _cheat_map(tryout_id)
    for a in attempts:
        a["student"] = smap.get(a["student_id"])
        a["cheat_total"] = cheat.get(a["student_id"], 0)
        a["tab_violations"] = a.get("tab_violations", 0)
    return sorted(attempts, key=lambda a: a.get("percentage", 0), reverse=True)


async def _cheat_map(tryout_id: str) -> dict:
    rows = await db.tryout_violations.find({"tryout_id": tryout_id}, {"_id": 0}).to_list(10000)
    return {r["student_id"]: r.get("total", 0) for r in rows}


@router.get("/tryouts/{tryout_id}/monitor")
async def tryout_monitor(tryout_id: str, user: dict = Depends(admin_only)):
    t = await db.tryouts.find_one({"id": tryout_id}, {"_id": 0})
    if not t:
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    total_q = await db.questions.count_documents({"tryout_id": tryout_id})
    attempts = await db.attempts.find({"tryout_id": tryout_id}, {"_id": 0, "per_question": 0}).to_list(5000)
    grants = await db.tryout_retakes.find({"tryout_id": tryout_id, "used": False}, {"_id": 0}).to_list(5000)
    sids = list({a["student_id"] for a in attempts} | {g["student_id"] for g in grants})
    smap = {u["id"]: u for u in await db.users.find({"id": {"$in": sids}}, {"_id": 0, "id": 1, "name": 1, "email": 1}).to_list(5000)}
    cheat = await _cheat_map(tryout_id)
    rows = []
    for a in attempts:
        answers = a.get("answers") or {}
        status = a["status"]
        left = remaining_seconds(a, t) if status == "in_progress" else 0
        if status == "in_progress" and left <= 0:
            status = "expired"
        rows.append({
            "attempt_id": a["id"], "student": smap.get(a["student_id"]), "student_id": a["student_id"],
            "status": status, "answered_count": a.get("answered_count", answered_count(answers)),
            "total_questions": total_q, "started_at": a.get("started_at"), "submitted_at": a.get("submitted_at"),
            "last_saved_at": a.get("last_saved_at") or (a.get("submitted_at") if a["status"] == "submitted" else None),
            "remaining_seconds": left, "percentage": a.get("percentage") if a["status"] == "submitted" else None,
            "retake": bool(a.get("bypass_schedule")),
            "tab_violations": a.get("tab_violations", 0), "cheat_total": cheat.get(a["student_id"], 0),
            "auto_submitted": a.get("auto_submit_reason") == "tab_violation",
        })
    for g in grants:
        rows.append({"attempt_id": None, "student": smap.get(g["student_id"]), "student_id": g["student_id"],
                     "status": "retake_pending", "answered_count": 0, "total_questions": total_q,
                     "started_at": None, "submitted_at": None, "last_saved_at": None, "remaining_seconds": 0,
                     "percentage": None, "retake": True, "granted_at": g.get("created_at"),
                     "tab_violations": 0, "cheat_total": cheat.get(g["student_id"], 0), "auto_submitted": False})
    order = {"in_progress": 0, "expired": 1, "retake_pending": 2, "submitted": 3}
    rows.sort(key=lambda r: (order.get(r["status"], 9), (r["student"] or {}).get("name") or ""))
    summary = {k: sum(1 for r in rows if r["status"] == k) for k in order}
    return {"tryout": {"id": t["id"], "title": t["title"], "start_at": t.get("start_at"), "end_at": t.get("end_at"),
                       "duration_minutes": t.get("duration_minutes"), "schedule_status": schedule_status(t)},
            "summary": summary, "rows": rows, "server_now": now_iso()}


@router.post("/tryouts/{tryout_id}/attempts/{attempt_id}/reset")
async def reset_attempt(tryout_id: str, attempt_id: str, background: BackgroundTasks, user: dict = Depends(admin_only)):
    a = await db.attempts.find_one({"id": attempt_id, "tryout_id": tryout_id}, {"_id": 0, "student_id": 1, "status": 1})
    if not a:
        raise HTTPException(status_code=404, detail="Data pengerjaan tidak ditemukan")
    sid = a["student_id"]
    await db.attempts.delete_many({"tryout_id": tryout_id, "student_id": sid})
    await db.tryout_retakes.delete_many({"tryout_id": tryout_id, "student_id": sid, "used": False})
    await db.tryout_retakes.insert_one({"id": new_id(), "tryout_id": tryout_id, "student_id": sid, "used": False,
                                        "granted_by": user["id"], "created_at": now_iso()})
    if a["status"] == "submitted":
        from routes_student import _recalibrate_tryout, align_tryout_irt
        await align_tryout_irt(tryout_id)
        background.add_task(_recalibrate_tryout, tryout_id)
    return {"ok": True}


@router.delete("/tryouts/{tryout_id}/retakes/{student_id}")
async def cancel_retake(tryout_id: str, student_id: str, user: dict = Depends(admin_only)):
    await db.tryout_retakes.delete_many({"tryout_id": tryout_id, "student_id": student_id, "used": False})
    return {"ok": True}


@router.post("/tryouts/{tryout_id}/attempts/{attempt_id}/force-submit")
async def force_submit(tryout_id: str, attempt_id: str, background: BackgroundTasks, user: dict = Depends(admin_only)):
    a = await db.attempts.find_one({"id": attempt_id, "tryout_id": tryout_id}, {"_id": 0})
    if not a:
        raise HTTPException(status_code=404, detail="Data pengerjaan tidak ditemukan")
    if a["status"] == "submitted":
        raise HTTPException(status_code=400, detail="Sudah dikumpulkan")
    t = await db.tryouts.find_one({"id": tryout_id}, {"_id": 0})
    questions = await db.questions.find({"tryout_id": tryout_id}, {"_id": 0}).to_list(1000)
    result = grade_attempt(questions, a.get("answers") or {})
    irt_score = irt.score_attempt(questions, result.get("per_question"), irt.resolve_scale(t))
    await db.attempts.update_one({"id": attempt_id}, {"$set": {
        "status": "submitted", "submitted_at": now_iso(), "force_submitted_by": user["id"], **result, **irt_score}})
    from routes_student import _recalibrate_tryout, align_tryout_irt
    await align_tryout_irt(tryout_id)
    background.add_task(_recalibrate_tryout, tryout_id)
    return {"ok": True, "percentage": result["percentage"]}


# ---------- Weakness Analysis & Score Export (all schools) ----------
@router.get("/analysis")
async def analysis_report(user: dict = Depends(admin_only)):
    return await analysis.build_report(None)


@router.get("/analysis/scores.csv")
async def analysis_scores_csv(user: dict = Depends(admin_only)):
    rep = await analysis.build_report(None)
    return Response(content=analysis.scores_csv(rep), media_type="text/csv",
                    headers={"Content-Disposition": "attachment; filename=nilai_siswa.csv"})


@router.get("/analysis/recap.csv")
async def analysis_recap_csv(user: dict = Depends(admin_only)):
    rep = await analysis.build_report(None)
    return Response(content=analysis.recap_csv(rep), media_type="text/csv",
                    headers={"Content-Disposition": "attachment; filename=rekap_kelemahan_siswa.csv"})


@router.get("/analysis/items.csv")
async def analysis_items_csv(user: dict = Depends(admin_only)):
    rep = await analysis.build_report(None)
    return Response(content=analysis.items_csv(rep), media_type="text/csv",
                    headers={"Content-Disposition": "attachment; filename=analisis_butir_soal.csv"})


# ---------- Broadcasts ----------
@router.get("/broadcasts")
async def list_broadcasts(user: dict = Depends(admin_only)):
    return await db.broadcasts.find({}, {"_id": 0}).sort("created_at", -1).to_list(200)


@router.post("/broadcasts")
async def create_broadcast(body: BroadcastBody, user: dict = Depends(admin_only)):
    doc = {"id": new_id(), **body.model_dump(), "author": user["name"], "created_at": now_iso()}
    await db.broadcasts.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


@router.delete("/broadcasts/{item_id}")
async def delete_broadcast(item_id: str, user: dict = Depends(admin_only)):
    await db.broadcasts.delete_one({"id": item_id})
    return {"ok": True}


# ---------- Users & Schools ----------
@router.get("/users")
async def list_users(role: Optional[str] = None, user: dict = Depends(admin_only)):
    q = {"role": role} if role else {}
    items = await db.users.find(q, {"_id": 0, "password_hash": 0}).sort("created_at", -1).to_list(500)
    schools = await db.schools.find({}, {"_id": 0}).to_list(200)
    smap = {s["id"]: s["name"] for s in schools}
    for u in items:
        u["school_name"] = smap.get(u.get("school_id"))
    return items


@router.post("/users")
async def create_user(body: UserBody, user: dict = Depends(admin_only)):
    if body.role not in ("tutor", "proctor", "student"):
        raise HTTPException(status_code=400, detail="Peran tidak valid")
    email = body.email.lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email sudah terdaftar")
    doc = {
        "id": new_id(),
        "email": email,
        "password_hash": hash_password(body.password),
        "name": body.name,
        "role": body.role,
        "status": "approved",
        "phone": body.phone,
        "school_id": body.school_id,
        "qualifications": body.qualifications,
        "picture": None,
        "auth_provider": "password",
        "created_at": now_iso(),
    }
    await db.users.insert_one(doc)
    return {"ok": True, "id": doc["id"]}


@router.put("/users/{user_id}")
async def update_user(user_id: str, body: UserUpdateBody, user: dict = Depends(admin_only)):
    target = await db.users.find_one({"id": user_id})
    if not target:
        raise HTTPException(status_code=404, detail="Pengguna tidak ditemukan")
    updates = {}
    if body.role is not None:
        if body.role not in ("admin", "tutor", "proctor", "student"):
            raise HTTPException(status_code=400, detail="Peran tidak valid")
        updates["role"] = body.role
    if body.status is not None:
        if body.status not in ("pending", "approved", "rejected"):
            raise HTTPException(status_code=400, detail="Status tidak valid")
        updates["status"] = body.status
    if body.school_id is not None:
        updates["school_id"] = body.school_id or None
    if not updates:
        raise HTTPException(status_code=400, detail="Tidak ada perubahan")
    await db.users.update_one({"id": user_id}, {"$set": updates})
    return {"ok": True, **updates}


@router.delete("/users/{user_id}")
async def delete_user(user_id: str, user: dict = Depends(admin_only)):
    if user_id == user["id"]:
        raise HTTPException(status_code=400, detail="Tidak dapat menghapus akun sendiri")
    await db.users.delete_one({"id": user_id})
    return {"ok": True}


@router.get("/schools")
async def list_schools(user: dict = Depends(admin_only)):
    items = await db.schools.find({}, {"_id": 0}).sort("name", 1).to_list(200)
    for s in items:
        s["student_count"] = await db.users.count_documents({"role": "student", "school_id": s["id"]})
    return items


@router.post("/schools")
async def create_school(body: SchoolBody, user: dict = Depends(admin_only)):
    doc = {"id": new_id(), **body.model_dump(), "created_at": now_iso()}
    await db.schools.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


# ---------- Course Content: Lessons & Exercises ----------
class LessonBody(BaseModel):
    title: str
    description: Optional[str] = None
    video_type: str = "youtube"  # youtube | upload
    video_url: Optional[str] = None
    attachments: List[dict] = []  # [{name, url, content_type}]
    order: int = 0


class ExerciseBody(BaseModel):
    title: str
    subject: Optional[str] = "Latihan"
    duration_minutes: int = 20
    description: Optional[str] = None


@router.get("/courses/{course_id}/content")
async def course_content(course_id: str, user: dict = Depends(admin_only)):
    course = await db.courses.find_one({"id": course_id}, {"_id": 0})
    if not course:
        raise HTTPException(status_code=404, detail="Kursus tidak ditemukan")
    lessons = await db.lessons.find({"course_id": course_id}, {"_id": 0}).sort("order", 1).to_list(200)
    exercises = await db.tryouts.find({"course_id": course_id, "kind": "exercise"}, {"_id": 0}).sort("created_at", 1).to_list(200)
    for ex in exercises:
        ex["question_count"] = await db.questions.count_documents({"tryout_id": ex["id"]})
        ex["attempt_count"] = await db.attempts.count_documents({"tryout_id": ex["id"], "status": "submitted"})
    return {"course": course, "lessons": lessons, "exercises": exercises}


@router.post("/courses/{course_id}/lessons")
async def create_lesson(course_id: str, body: LessonBody, user: dict = Depends(admin_only)):
    if not await db.courses.find_one({"id": course_id}):
        raise HTTPException(status_code=404, detail="Kursus tidak ditemukan")
    count = await db.lessons.count_documents({"course_id": course_id})
    doc = {"id": new_id(), "course_id": course_id, **body.model_dump(), "created_at": now_iso()}
    if not doc.get("order"):
        doc["order"] = count + 1
    await db.lessons.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


@router.delete("/lessons/{lesson_id}")
async def delete_lesson(lesson_id: str, user: dict = Depends(admin_only)):
    await db.lessons.delete_one({"id": lesson_id})
    return {"ok": True}


@router.post("/courses/{course_id}/exercises")
async def create_exercise(course_id: str, body: ExerciseBody, user: dict = Depends(admin_only)):
    if not await db.courses.find_one({"id": course_id}):
        raise HTTPException(status_code=404, detail="Kursus tidak ditemukan")
    doc = {
        "id": new_id(), "title": body.title, "description": body.description,
        "subject": body.subject or "Latihan", "duration_minutes": body.duration_minutes,
        "start_at": now_iso(), "end_at": None, "published": True,
        "course_id": course_id, "kind": "exercise",
        "created_by": user["id"], "created_at": now_iso(),
    }
    await db.tryouts.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


# ---------- Bulk Import Questions (CSV / Excel) ----------
_LETTERS = ["a", "b", "c", "d", "e", "f"]


def _row_to_question(row: dict, order: int):
    t = str(row.get("type") or "").strip().lower()
    t = {"bs_majemuk": "truefalse_multi", "benar_salah_majemuk": "truefalse_multi", "truefalse-multi": "truefalse_multi"}.get(t, t)
    if t not in ("single", "multiple", "truefalse", "truefalse_multi", "essay"):
        return None, f"Tipe tidak valid: {t!r}"
    text = str(row.get("text") or "").strip()
    if not text:
        return None, "Kolom 'text' kosong"
    try:
        points = int(float(str(row.get("points") or "").strip()))
    except Exception:
        return None, "Kolom 'points' wajib diisi angka"
    if points <= 0:
        return None, "Kolom 'points' harus lebih dari 0"
    options, correct = [], []
    for i, l in enumerate(_LETTERS):
        v = row.get(f"option_{l}")
        v = str(v).strip() if v is not None else ""
        if v:
            options.append({"id": f"o{i + 1}", "text": v})
    correct_raw = str(row.get("correct") or "").strip()
    if t in ("single", "multiple"):
        for tok in re.split(r"[;,|]", correct_raw):
            tok = tok.strip().lower()
            if not tok:
                continue
            if tok in _LETTERS:
                correct.append(f"o{_LETTERS.index(tok) + 1}")
            else:
                for o in options:
                    if o["text"].strip().lower() == tok:
                        correct.append(o["id"])
        if not options or not correct:
            return None, "Opsi atau kunci tidak lengkap"
    elif t == "truefalse":
        val = correct_raw.strip().lower()
        if val in ("true", "benar", "b", "ya", "1"):
            val = "true"
        elif val in ("false", "salah", "s", "tidak", "0"):
            val = "false"
        else:
            return None, "Kunci benar/salah tidak valid"
        correct, options = [val], []
    elif t == "truefalse_multi":
        keys = [k.strip().lower() for k in re.split(r"[;,|]", correct_raw) if k.strip()]
        if not options or len(keys) != len(options):
            return None, "Jumlah kunci (mis. B;S;B) harus sama dengan jumlah pernyataan di option_a.."
        for o, k in zip(options, keys):
            if k in ("true", "benar", "b", "ya", "1"):
                correct.append(f"{o['id']}:true")
            elif k in ("false", "salah", "s", "tidak", "0"):
                correct.append(f"{o['id']}:false")
            else:
                return None, f"Kunci benar/salah tidak valid: {k!r}"
    else:  # essay
        correct = [c.strip() for c in re.split(r"[|;]", correct_raw) if c.strip()]
        options = []
        if not correct:
            return None, "Kunci esai kosong"
    comp = str(row.get("competency") or row.get("kompetensi") or "").strip().lower()
    if comp not in ("numerasi", "literasi", "umum"):
        return None, "Kolom 'competency' wajib diisi: numerasi, literasi, atau umum"
    dl = str(row.get("difficulty_label") or row.get("tingkat") or row.get("kesulitan") or "").strip().lower()
    _MAP = {"mudah": "mudah", "easy": "mudah", "sedang": "sedang", "medium": "sedang",
            "sulit": "sulit", "susah": "sulit", "hard": "sulit"}
    dl = _MAP.get(dl)
    if not dl:
        return None, "Kolom 'tingkat' wajib diisi: mudah, sedang, atau sulit"
    lt = str(row.get("label_benar") or "").strip()
    lf = str(row.get("label_salah") or "").strip()
    tf_labels = [lt or "Benar", lf or "Salah"] if t in ("truefalse", "truefalse_multi") and (lt or lf) else None
    return {"type": t, "text": text, "options": options, "correct_answers": correct, "tf_labels": tf_labels,
            "points": points, "order": order, "competency": comp, "difficulty_label": dl,
            "bidang": _clean_text(str(row.get("bidang") or row.get("mapel") or "")),
            "materi": _clean_text(str(row.get("materi") or row.get("topik") or "")),
            "explanation": _clean_text(str(row.get("pembahasan") or row.get("explanation") or ""))}, None


def _parse_sheet(data: bytes, fname: str):
    rows = []
    if fname.endswith(".xlsx") or fname.endswith(".xls"):
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        headers = None
        for r in ws.iter_rows(values_only=True):
            if headers is None:
                headers = [str(c).strip().lower() if c is not None else "" for c in r]
                continue
            if all(c is None or str(c).strip() == "" for c in r):
                continue
            rows.append({headers[i]: r[i] for i in range(len(headers)) if i < len(r)})
    else:
        text = data.decode("utf-8-sig", errors="ignore")
        for r in csv.DictReader(io.StringIO(text)):
            rows.append({(k or "").strip().lower(): v for k, v in r.items()})
    return rows


def _parse_zip(data: bytes):
    images, sheet, sheet_name = {}, None, ""
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for info in zf.infolist():
            name = info.filename
            base = name.replace("\\", "/").split("/")[-1]
            if info.is_dir() or not base or name.startswith("__MACOSX") or base.startswith("."):
                continue
            low = base.lower()
            if low.endswith((".xlsx", ".xls", ".csv")) and sheet is None:
                sheet, sheet_name = zf.read(info), low
            elif low.rsplit(".", 1)[-1] in IMAGE_EXT:
                images[low] = (base, zf.read(info))
    if sheet is None:
        raise ValueError("ZIP harus berisi 1 file Excel (.xlsx) atau CSV")
    return _parse_sheet(sheet, sheet_name), images


REQUIRED_COLUMNS = ["type", "text", "option_a", "option_b", "option_c", "option_d", "option_e",
                    "correct", "points", "competency", "tingkat"]


def _dup_key(text, options=None) -> str:
    norm = lambda v: re.sub(r"\s+", " ", str(v or "")).strip().lower()
    opts = sorted(norm(o.get("text") if isinstance(o, dict) else o) for o in (options or []))
    return norm(text) + "||" + "|".join(o for o in opts if o)


async def _existing_keys(tryout_id):
    if not tryout_id:
        return set()
    qs = await db.questions.find({"tryout_id": tryout_id}, {"_id": 0, "text": 1, "options": 1}).to_list(5000)
    return {_dup_key(q.get("text"), q.get("options")) for q in qs}


def _missing_columns(rows):
    have = set()
    for r in rows:
        have |= set(r.keys())
    return [c for c in REQUIRED_COLUMNS if c not in have]


_IMG_FIELDS = (("gambar", ("gambar", "image", "image_url"), "gambar soal"),
               ("gambar_pembahasan", ("gambar_pembahasan", "explanation_image", "explanation_image_url"), "gambar pembahasan"))


def _is_url(v: str) -> bool:
    return v.startswith(("https://", "http://", "/api/"))


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _normalize_rows(rows):
    out = []
    for r in rows:
        n = {k: _cell(v) for k, v in r.items() if k}
        for key, aliases, _ in _IMG_FIELDS:
            n[key] = next((n[x] for x in aliases if n.get(x)), "")
        out.append(n)
    return out


async def _upload_row_images(rows, images, user_id):
    """Unggah gambar dari ZIP & ganti nama file dengan URL. Return {index: [peringatan]}."""
    warns = {}
    for i, r in enumerate(rows):
        for key, _, label in _IMG_FIELDS:
            v = r.get(key) or ""
            if not v or _is_url(v):
                continue
            hit = images.get(v.split("/")[-1].lower())
            if not hit:
                continue
            try:
                r[key] = await save_question_image(hit[0], hit[1], user_id)
            except Exception as e:
                warns.setdefault(i, []).append(f"{label} {v!r} gagal diunggah ({e})")
    return warns


def _validate_row(row: dict, order: int = 0):
    q, err = _row_to_question(row, order)
    warns = []
    for key, _, label in _IMG_FIELDS:
        v = row.get(key) or ""
        if v and not _is_url(v):
            warns.append(f"{label} {v!r} tidak ditemukan (unggah ZIP berisi Excel + gambar); soal diimpor tanpa gambar itu")
    if q:
        q["image_url"] = row["gambar"] if _is_url(row.get("gambar") or "") else None
        q["explanation_image_url"] = row["gambar_pembahasan"] if _is_url(row.get("gambar_pembahasan") or "") else None
    return q, err, warns


def _preview_payload(rows, existing, extra_warns=None):
    out, seen = [], {}
    for i, r in enumerate(rows):
        q, err, warns = _validate_row(r)
        dup = None
        if q:
            k = _dup_key(q["text"], q["options"])
            if k in existing:
                dup = "Soal yang sama sudah ada di Try Out ini, akan dilewati"
            elif k in seen:
                dup = f"Sama dengan baris {seen[k] + 2} di berkas ini, akan dilewati"
            else:
                seen[k] = i
        out.append({"raw": r, "ok": err is None and dup is None, "error": err, "duplicate": dup,
                    "warnings": (extra_warns or {}).get(i, []) + warns, "question": q})
    return {"rows": out, "valid": sum(1 for x in out if x["ok"]),
            "invalid": sum(1 for x in out if x["error"]), "duplicates": sum(1 for x in out if x["duplicate"])}


async def _read_upload(file: UploadFile, user_id: str):
    data = await file.read()
    fname = (file.filename or "").lower()
    images = {}
    try:
        if fname.endswith(".zip"):
            rows, images = _parse_zip(data)
        else:
            rows = _parse_sheet(data, fname)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Gagal membaca berkas: {e}")
    rows = _normalize_rows(rows)
    missing = _missing_columns(rows)
    if rows and missing:
        raise HTTPException(status_code=400, detail=f"Kolom wajib tidak ada di berkas: {', '.join(missing)}. Gunakan template terbaru.")
    warns = await _upload_row_images(rows, images, user_id)
    return rows, warns


async def _insert_rows(tryout_id: str, rows):
    tryout = await db.tryouts.find_one({"id": tryout_id}, {"_id": 0, "subject": 1})
    if not tryout:
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    start = await db.questions.count_documents({"tryout_id": tryout_id})
    existing = await _existing_keys(tryout_id)
    imported, errors, duplicates = 0, [], []
    for idx, row in enumerate(rows):
        q, err, warns = _validate_row(row, start + imported + 1)
        if err:
            errors.append(f"Baris {idx + 2}: {err}")
            continue
        k = _dup_key(q["text"], q["options"])
        if k in existing:
            duplicates.append(f"Baris {idx + 2}: \"{q['text'][:60]}\" sudah ada")
            continue
        existing.add(k)
        errors += [f"Baris {idx + 2}: {w}" for w in warns]
        q.update({"id": new_id(), "tryout_id": tryout_id,
                  "bidang": q.get("bidang") or _clean_text(tryout.get("subject")),
                  "irt_a": None, "irt_b": None, "irt_c": None, "irt_calibrated": False,
                  "p_value": None, "point_biserial": None, "response_count": 0})
        await db.questions.insert_one(q)
        imported += 1
    return {"imported": imported, "errors": errors, "duplicates": duplicates}


class RowsBody(BaseModel):
    rows: List[dict]


@router.post("/tryouts/{tryout_id}/questions/import")
async def import_questions(tryout_id: str, file: UploadFile = File(...), user: dict = Depends(admin_only)):
    if not await db.tryouts.find_one({"id": tryout_id}):
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    rows, _ = await _read_upload(file, user["id"])
    return await _insert_rows(tryout_id, rows)


@router.post("/questions/import/preview")
async def import_preview(file: UploadFile = File(...), tryout_id: Optional[str] = None, user: dict = Depends(admin_only)):
    rows, warns = await _read_upload(file, user["id"])
    if not rows:
        raise HTTPException(status_code=400, detail="Berkas tidak berisi baris soal")
    return _preview_payload(rows, await _existing_keys(tryout_id), warns)


@router.post("/questions/import/validate")
async def import_validate(body: RowsBody, tryout_id: Optional[str] = None, user: dict = Depends(admin_only)):
    return _preview_payload(_normalize_rows(body.rows), await _existing_keys(tryout_id))


class IdsBody(BaseModel):
    question_ids: List[str]


@router.get("/question-bank/duplicates")
async def question_duplicates(user: dict = Depends(admin_only)):
    qs = await db.questions.find({}, {"_id": 0, "id": 1, "tryout_id": 1, "type": 1, "text": 1, "options": 1, "order": 1,
                                       "bidang": 1, "materi": 1, "explanation": 1}).to_list(50000)
    groups = {}
    for q in qs:
        groups.setdefault(_dup_key(q.get("text"), q.get("options")), []).append(q)
    groups = [g for g in groups.values() if len(g) > 1]
    tids = list({q["tryout_id"] for g in groups for q in g})
    tmap = {t["id"]: t for t in await db.tryouts.find({"id": {"$in": tids}}, {"_id": 0, "id": 1, "title": 1, "kind": 1}).to_list(5000)}
    used = {r["_id"] for r in await db.attempts.aggregate([
        {"$match": {"tryout_id": {"$in": tids}, "status": "submitted"}}, {"$group": {"_id": "$tryout_id"}}]).to_list(5000)}
    out = []
    for g in groups:
        items = [{"id": q["id"], "tryout_id": q["tryout_id"], "tryout_title": (tmap.get(q["tryout_id"]) or {}).get("title", "-"),
                  "type": q.get("type"), "order": q.get("order"), "bidang": q.get("bidang"), "materi": q.get("materi"),
                  "has_explanation": bool(q.get("explanation")), "tryout_has_results": q["tryout_id"] in used} for q in g]
        per_t = {}
        for it in items:
            per_t.setdefault(it["tryout_id"], []).append(it)
        out.append({"text": g[0].get("text", ""), "count": len(g), "items": items,
                    "same_tryout": any(len(v) > 1 for v in per_t.values())})
    out.sort(key=lambda x: (not x["same_tryout"], -x["count"]))
    return {"groups": out, "total_groups": len(out), "same_tryout_groups": sum(1 for x in out if x["same_tryout"])}


@router.post("/questions/bulk-delete")
async def bulk_delete_questions(body: IdsBody, user: dict = Depends(admin_only)):
    if not body.question_ids:
        raise HTTPException(status_code=400, detail="Pilih soal yang akan dihapus")
    res = await db.questions.delete_many({"id": {"$in": body.question_ids}})
    return {"deleted": res.deleted_count}


@router.post("/tryouts/{tryout_id}/questions/import-rows")
async def import_rows(tryout_id: str, body: RowsBody, user: dict = Depends(admin_only)):
    return await _insert_rows(tryout_id, _normalize_rows(body.rows))


_TPL_HEADERS = ["type", "text", "option_a", "option_b", "option_c", "option_d", "option_e", "correct", "points",
                "competency", "tingkat", "bidang", "materi", "pembahasan", "gambar", "gambar_pembahasan", "label_benar", "label_salah"]
_TPL_SAMPLES = [
    {"type": "single", "text": "Berapa hasil 2+2?", "option_a": "3", "option_b": "4", "option_c": "5", "option_d": "6",
     "correct": "B", "points": 10, "competency": "numerasi", "tingkat": "mudah", "bidang": "Matematika", "materi": "Bilangan",
     "pembahasan": "2 + 2 = 4, jadi jawaban yang tepat adalah B."},
    {"type": "single", "text": r"Hasil dari $\frac{1}{2} + \frac{1}{3}$ adalah ...", "option_a": r"$\frac{5}{6}$",
     "option_b": r"$\frac{2}{5}$", "option_c": r"$\frac{1}{6}$", "option_d": r"$\frac{3}{5}$", "correct": "A", "points": 10,
     "competency": "numerasi", "tingkat": "sedang", "bidang": "Matematika", "materi": "Pecahan",
     "pembahasan": r"Samakan penyebut: $\frac{1}{2} + \frac{1}{3} = \frac{3}{6} + \frac{2}{6} = \frac{5}{6}$."},
    {"type": "single", "text": r"Akar-akar dari $$x^2 - 5x + 6 = 0$$ adalah ...", "option_a": r"$x=2$ dan $x=3$",
     "option_b": r"$x=-2$ dan $x=-3$", "option_c": r"$x=1$ dan $x=6$", "option_d": r"$x=-1$ dan $x=6$", "correct": "A",
     "points": 10, "competency": "numerasi", "tingkat": "sulit", "bidang": "Matematika", "materi": "Persamaan Kuadrat",
     "pembahasan": r"Faktorkan: $(x-2)(x-3)=0$, sehingga $x=2$ atau $x=3$."},
    {"type": "single", "text": "Perhatikan gambar. Luas segitiga tersebut adalah ...", "option_a": r"$12\text{ cm}^2$",
     "option_b": r"$24\text{ cm}^2$", "option_c": r"$6\text{ cm}^2$", "option_d": r"$18\text{ cm}^2$", "correct": "A",
     "points": 10, "competency": "numerasi", "tingkat": "sedang", "bidang": "Matematika", "materi": "Bangun Datar",
     "pembahasan": r"Luas segitiga $= \frac{1}{2} \times alas \times tinggi = \frac{1}{2} \times 6 \times 4 = 12\text{ cm}^2$.",
     "gambar": "segitiga.png", "gambar_pembahasan": "pembahasan_segitiga.png"},
    {"type": "multiple", "text": "Pilih bilangan genap", "option_a": "2", "option_b": "3", "option_c": "4", "option_d": "5",
     "correct": "A;C", "points": 10, "competency": "numerasi", "tingkat": "sedang", "bidang": "Matematika", "materi": "Bilangan",
     "pembahasan": "Bilangan genap habis dibagi 2: yaitu 2 dan 4."},
    {"type": "truefalse", "text": "Bumi berbentuk bulat", "correct": "benar", "points": 10, "competency": "literasi",
     "tingkat": "mudah", "bidang": "IPA", "materi": "Tata Surya", "pembahasan": "Bumi berbentuk bulat (sedikit pepat di kutub)."},
    {"type": "truefalse_multi", "text": "Bacalah teks berikut. Rani membeli 3 buku seharga Rp5.000 per buku dan membayar Rp20.000. Tentukan benar atau salah setiap pernyataan.",
     "option_a": "Total harga buku Rp15.000", "option_b": "Uang kembalian Rp10.000", "option_c": "Jika membeli 4 buku uangnya pas",
     "correct": "B;S;B", "points": 15, "competency": "numerasi", "tingkat": "sedang", "bidang": "Matematika", "materi": "Aritmetika Sosial",
     "pembahasan": "3 x Rp5.000 = Rp15.000 (benar). Kembalian Rp20.000 - Rp15.000 = Rp5.000 (salah). 4 x Rp5.000 = Rp20.000 (benar).",
     "label_benar": "Sesuai", "label_salah": "Tidak Sesuai"},
    {"type": "essay", "text": "Ibu kota Indonesia?", "correct": "jakarta|dki jakarta", "points": 10, "competency": "literasi",
     "tingkat": "sulit", "bidang": "IPS", "materi": "Geografi Indonesia", "pembahasan": "Ibu kota Indonesia adalah Jakarta (DKI Jakarta)."},
]


@router.get("/questions/template")
async def questions_template(user: dict = Depends(admin_only)):
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=_TPL_HEADERS, extrasaction="ignore")
    w.writeheader()
    for row in _TPL_SAMPLES:
        w.writerow(row)
    return Response(content="\ufeff" + buf.getvalue(), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": "attachment; filename=template_soal.csv"})


@router.get("/questions/template.xlsx")
async def questions_template_xlsx(user: dict = Depends(admin_only)):
    """Template Excel untuk pembuatan soal skala besar (banyak soal sekaligus)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Soal"
    ws.append(_TPL_HEADERS)
    for row in _TPL_SAMPLES:
        ws.append([row.get(h, "") for h in _TPL_HEADERS])
    from openpyxl.styles import PatternFill, Font
    for cell in ws[1]:
        req = cell.value in REQUIRED_COLUMNS
        cell.fill = PatternFill("solid", fgColor="0E7490" if req else "E2E8F0")
        cell.font = Font(bold=True, color="FFFFFF" if req else "475569")
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(max(width, 10), 48)

    info = wb.create_sheet("Petunjuk")
    guide = [
        ["Petunjuk Pengisian Template Soal (Impor Massal)"],
        [""],
        ["WAJIB", "Kolom header biru wajib ada: " + ", ".join(REQUIRED_COLUMNS) + ". Isi per baris: type, text, correct, points, competency, tingkat (opsi sesuai tipe soal)."],
        ["OPSIONAL", "Kolom header abu-abu (bidang, materi, pembahasan, gambar, gambar_pembahasan, label_benar, label_salah) boleh dikosongkan & diatur nanti di halaman admin."],
        ["Duplikat", "Soal yang sama (teks + pilihan) dengan soal yang sudah ada di Try Out tujuan, atau berulang di berkas, otomatis dilewati dan diberi pemberitahuan."],
        [""],
        ["Kolom", "Keterangan"],
        ["type", "Tipe soal: single (PG), multiple (PG kompleks), truefalse (1 pernyataan), truefalse_multi (Benar/Salah majemuk: 1 teks + beberapa pernyataan), essay"],
        ["text", "Teks pertanyaan"],
        ["option_a..e", "Pilihan jawaban (untuk single/multiple). Kosongkan untuk truefalse/essay"],
        ["correct", "Kunci: huruf A-E (pisah ; untuk multiple), benar/salah untuk truefalse, B;S;B (urut sesuai option_a..e) untuk truefalse_multi, atau teks (pisah |) untuk essay"],
        ["truefalse_multi", "Tulis teks deskripsi/stimulus di kolom text, lalu tiap pernyataan di option_a, option_b, dst. Nilai dihitung proporsional dari jumlah pernyataan yang benar."],
        ["label_benar / label_salah", "OPSIONAL untuk truefalse & truefalse_multi. Ganti teks tombol, mis. Sesuai / Tidak Sesuai, Fakta / Opini, Ya / Tidak. Kosong = Benar / Salah."],
        ["points", "Poin/bobot soal (angka)"],
        ["competency", "numerasi | literasi | umum (untuk laporan AKM)"],
        ["tingkat", "Tingkat kesulitan IRT: mudah | sedang | sulit"],
        ["bidang", "OPSIONAL. Bidang / mata pelajaran soal (mis. Matematika). Kosong = otomatis mata pelajaran Try Out tujuan impor."],
        ["materi", "OPSIONAL. Materi / topik soal (mis. Pecahan, Aljabar). Dipakai untuk filter Bank Soal."],
        ["gambar_pembahasan", "OPSIONAL. Nama file gambar untuk pembahasan (mis. pembahasan_segitiga.png) di dalam ZIP yang sama, atau link https://..."],
        ["pembahasan", "OPSIONAL. Penjelasan lengkap jawaban, tampil ke siswa jika admin membuka pembahasan. Mendukung rumus $...$."],
        ["gambar", "OPSIONAL. Isi nama file gambar (mis. segitiga.png) lalu unggah ZIP berisi file Excel ini + semua gambarnya. Bisa juga diisi link https://... Kosongkan bila soal tanpa gambar."],
        [""],
        ["Rumus Matematika (LaTeX)", "Tulis rumus di antara tanda $...$ (sebaris) atau $$...$$ (baris sendiri). Berlaku di kolom text & option."],
        ["Contoh pecahan", r"$\frac{a}{b}$"],
        ["Contoh pangkat & akar", r"$x^2 + \sqrt{y} = \sqrt[3]{27}$"],
        ["Contoh integral & limit", r"$$\int_0^1 x^2\,dx = \frac{1}{3}$$   $\lim_{x \to 0} \frac{\sin x}{x} = 1$"],
        ["Contoh simbol", r"$\alpha, \beta, \pi, \theta, \le, \ge, \ne, \times, \div, \pm, \infty$"],
        ["Contoh matriks", r"$\begin{pmatrix} 1 & 2 \\ 3 & 4 \end{pmatrix}$"],
        ["Tanda dolar biasa", r"Tulis \$ untuk menampilkan simbol $ (mis. harga \$5)"],
        [""],
        ["Rekomendasi komposisi IRT", "±30% Mudah, 40% Sedang, 30% Sulit"],
        ["Catatan", "Parameter IRT (a,b,c) akan dikalibrasi otomatis dari hasil jawaban peserta."],
    ]
    for row in guide:
        info.append(row)
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 100

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=template_soal.xlsx"},
    )


# ---------- Bulk Import Course Materials (ZIP) ----------
_VIDEO_EXT = {"mp4", "webm", "mov", "m4v"}
_DOC_EXT = {"pdf", "doc", "docx", "ppt", "pptx"}


@router.post("/courses/{course_id}/lessons/import-zip")
async def import_zip(course_id: str, file: UploadFile = File(...), user: dict = Depends(admin_only)):
    if not await db.courses.find_one({"id": course_id}):
        raise HTTPException(status_code=404, detail="Kursus tidak ditemukan")
    data = await file.read()
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except Exception:
        raise HTTPException(status_code=400, detail="Berkas ZIP tidak valid")

    created, skipped, errors = 0, 0, []
    order = await db.lessons.count_documents({"course_id": course_id})
    for name in zf.namelist():
        base = name.split("/")[-1]
        if not base or base.startswith(".") or "__MACOSX" in name:
            continue
        ext = base.rsplit(".", 1)[-1].lower() if "." in base else ""
        if ext not in _VIDEO_EXT and ext not in _DOC_EXT:
            skipped += 1
            continue
        try:
            content = zf.read(name)
            path = f"{APP_NAME}/uploads/{user['id']}/{new_id()}.{ext}"
            ct = MIME_TYPES.get(ext, "application/octet-stream")
            result = await asyncio.to_thread(put_object, path, content, ct)
        except Exception as e:
            errors.append(f"{base}: {e}")
            continue
        stored_path = result["path"]
        await db.files.insert_one({
            "id": new_id(), "storage_path": stored_path, "original_filename": base,
            "content_type": ct, "size": result.get("size", len(content)),
            "uploaded_by": user["id"], "is_deleted": False, "created_at": now_iso(),
        })
        url = f"/api/files/{stored_path}"
        order += 1
        title = base.rsplit(".", 1)[0]
        if ext in _VIDEO_EXT:
            doc = {"id": new_id(), "course_id": course_id, "title": title, "description": "Diimpor dari ZIP",
                   "video_type": "upload", "video_url": url, "attachments": [], "order": order, "created_at": now_iso()}
        else:
            doc = {"id": new_id(), "course_id": course_id, "title": title, "description": "Diimpor dari ZIP",
                   "video_type": "document", "video_url": None,
                   "attachments": [{"name": base, "url": url, "content_type": ct}], "order": order, "created_at": now_iso()}
        await db.lessons.insert_one(doc)
        created += 1
    return {"created": created, "skipped": skipped, "errors": errors}
